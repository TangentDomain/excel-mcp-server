#!/usr/bin/env python3
"""
测试openpyxl兼容性问题的处理
"""

import os
import sys
import tempfile
from pathlib import Path

# 添加项目根目录到Python路径
project_root = Path(__file__).parent
sys.path.insert(0, str(project_root))

try:
    from src.core.excel_search import ExcelSearcher
    print("✅ 成功导入ExcelSearcher")
except ImportError as e:
    print(f"❌ 导入ExcelSearcher失败: {e}")
    sys.exit(1)

def create_problematic_excel_file():
    """创建一个可能导致兼容性问题的Excel文件"""
    try:
        from openpyxl import Workbook
        from openpyxl.workbook.defined_name import DefinedName
        
        # 创建临时文件
        temp_file = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        temp_file.close()
        
        wb = Workbook()
        ws = wb.active
        ws.title = "测试数据"
        
        # 添加一些数据
        ws['A1'] = "姓名"
        ws['B1'] = "年龄"
        ws['A2'] = "张三"
        ws['B2'] = 25
        ws['A3'] = "李四"
        ws['B3'] = 30
        
        # 尝试添加命名范围（可能导致版本兼容性问题）
        try:
            # 创建命名范围，可能在某些openpyxl版本中有问题
            defined_name = DefinedName('TestRange', attr_text='测试数据!$A$1:$B$3')
            wb.defined_names.append(defined_name)
            print("✅ 成功添加命名范围")
        except Exception as e:
            print(f"⚠️  添加命名范围时出现问题 (预期的): {e}")
        
        wb.save(temp_file.name)
        wb.close()
        
        print(f"✅ 创建测试文件: {temp_file.name}")
        return temp_file.name
        
    except Exception as e:
        print(f"❌ 创建测试文件失败: {e}")
        return None

def test_compatibility_handling():
    """测试兼容性问题处理"""
    
    print("🔧 正在创建可能有问题的Excel文件...")
    test_file = create_problematic_excel_file()
    
    if not test_file:
        print("❌ 无法创建测试文件")
        return
    
    try:
        print("\n🔍 测试单文件搜索...")
        
        # 测试单文件搜索
        searcher = ExcelSearcher(test_file)
        result = searcher.regex_search(r'\d+', "", True, False)
        
        if result.success:
            print(f"✅ 单文件搜索成功! 找到 {len(result.data)} 个匹配项")
            if result.data:
                print(f"   示例匹配: {result.data[0].__dict__ if hasattr(result.data[0], '__dict__') else result.data[0]}")
        else:
            print(f"❌ 单文件搜索失败: {result.error}")
        
        # 测试目录搜索
        print("\n🗂️  测试目录搜索...")
        temp_dir = os.path.dirname(test_file)
        dir_result = ExcelSearcher.search_directory_static(
            temp_dir, r'\d+', "", True, False, False, ['.xlsx'], None, 10
        )
        
        if dir_result.success:
            print(f"✅ 目录搜索成功! 找到 {dir_result.metadata['total_matches']} 个匹配项")
            print(f"   搜索文件数: {dir_result.metadata['total_files_found']}")
            print(f"   成功文件: {len(dir_result.metadata['searched_files'])}")
            print(f"   跳过文件: {len(dir_result.metadata['skipped_files'])}")
            print(f"   错误文件: {len(dir_result.metadata['file_errors'])}")
            
            if dir_result.metadata['file_errors']:
                print("   文件错误详情:")
                for error in dir_result.metadata['file_errors']:
                    print(f"     - {error['file_path']}: {error['error']}")
        else:
            print(f"❌ 目录搜索失败: {dir_result.error}")
        
    except Exception as e:
        print(f"❌ 测试发生异常: {e}")
        import traceback
        traceback.print_exc()
        
    finally:
        # 清理临时文件
        try:
            if test_file and os.path.exists(test_file):
                os.unlink(test_file)
                print(f"🧹 已清理临时文件: {test_file}")
        except Exception as e:
            print(f"⚠️  清理临时文件失败: {e}")

if __name__ == "__main__":
    print("🧪 开始测试openpyxl兼容性处理")
    test_compatibility_handling()
    print("✨ 测试完成")
