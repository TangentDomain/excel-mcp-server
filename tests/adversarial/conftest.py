"""差分对抗测试共享 fixture 和辅助函数。

核心思路：SQLite 是真值来源（oracle）。
对同一个 Excel 文件，同时在 ExcelMCP 和 SQLite 上执行相同操作，
对比结果来发现 ExcelMCP 的 bug。
"""

from __future__ import annotations

import math
import os
import random
import shutil
import sqlite3
import uuid
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook

import excel_mcp_server_fastmcp.api.advanced_sql_query as _query_module
from excel_mcp_server_fastmcp.api.advanced_sql_query import (
    execute_advanced_delete_query,
    execute_advanced_insert_query,
    execute_advanced_sql_query,
    execute_advanced_update_query,
)

# ============================================================
# 种子数据定义
# ============================================================

SEED_TABLE = "商品"
SEED_COLUMNS = ["ID", "Name", "Price", "Stock", "Active"]
SEED_DATA = [
    [1, "铁剑", 100.0, 50, "是"],
    [2, "银剑", 200.0, 30, "是"],
    [3, "木盾", 50.0, 100, "否"],
    [4, "铁甲", 180.0, 20, "是"],
    [5, "皮甲", 80.0, 60, "是"],
    [6, "神器", 999.99, 1, "是"],
]


# ============================================================
# Excel 文件创建
# ============================================================


def _make_seed_wb() -> Workbook:
    """创建种子 Excel：单行表头 + 6 行数据"""
    wb = Workbook()
    ws = wb.active
    ws.title = SEED_TABLE
    ws.append(SEED_COLUMNS)
    for row in SEED_DATA:
        ws.append(row)
    return wb


def _save_wb(wb: Workbook, tmp_dir: Path, name: str) -> str:
    """保存 Workbook 到临时目录，返回路径"""
    path = tmp_dir / name
    wb.save(str(path))
    return str(path)


# ============================================================
# SQLite Oracle 创建
# ============================================================


def _create_oracle_db() -> sqlite3.Connection:
    """创建内存 SQLite oracle，schema 与种子 Excel 一致。

    注意：不使用 cmd_import（它添加 _rowid_ 列），
    直接建表保持与 ExcelMCP 的 SQL 引擎相同结构。
    """
    conn = sqlite3.connect(":memory:")
    col_defs = ", ".join(
        f"[{col}] {'REAL' if col == 'Price' else 'TEXT' if col == 'Name' or col == 'Active' else 'INTEGER'}"
        for col in SEED_COLUMNS
    )
    conn.execute(f"CREATE TABLE [{SEED_TABLE}] ({col_defs})")
    placeholders = ", ".join(["?"] * len(SEED_COLUMNS))
    for row in SEED_DATA:
        conn.execute(f"INSERT INTO [{SEED_TABLE}] VALUES ({placeholders})", row)
    conn.commit()
    return conn


# ============================================================
# 辅助函数：查询和写入两边
# ============================================================


def query_excel(excel_path: str, sql: str) -> dict[str, Any]:
    """在 ExcelMCP 上执行查询"""
    return execute_advanced_sql_query(excel_path, sql)


def query_sqlite(conn: sqlite3.Connection, sql: str) -> list[tuple]:
    """在 SQLite oracle 上执行查询，返回行列表"""
    cursor = conn.execute(sql)
    rows = cursor.fetchall()
    cursor.close()
    return rows


def write_excel(excel_path: str, sql: str, op_type: str = "update") -> dict[str, Any]:
    """在 ExcelMCP 上执行写操作"""
    if op_type == "update":
        return execute_advanced_update_query(excel_path, sql)
    elif op_type == "insert":
        return execute_advanced_insert_query(excel_path, sql)
    elif op_type == "delete":
        return execute_advanced_delete_query(excel_path, sql)
    else:
        raise ValueError(f"Unknown op_type: {op_type}")


def write_sqlite(conn: sqlite3.Connection, sql: str) -> int:
    """在 SQLite oracle 上执行写操作，返回 affected rows"""
    cursor = conn.execute(sql)
    affected = cursor.rowcount
    conn.commit()
    cursor.close()
    return affected


def reset_excel_engine():
    """重置 ExcelMCP 共享引擎缓存，避免跨测试污染"""
    _query_module._shared_engine = None


# ============================================================
# 差分对比函数
# ============================================================


def _normalize_cell(v: Any) -> Any:
    """归一化单元格值，用于比较。

    Excel 无法区分空字符串 '' 和 None（openpyxl 往返后 '' 变 None），
    所以在差分测试中将 '' 和 None 视为等价（这是 Excel 存储层的固有限制）。
    """
    if v is None:
        return None
    if isinstance(v, str) and v == "":
        return None  # '' 归一化为 None（Excel 存储限制）
    if isinstance(v, float) and math.isnan(v):
        return None
    if isinstance(v, float) and v == int(v):
        # ExcelMCP 可能返回 int，SQLite 可能返回 float（或反之）
        # 保留原类型，在对比时处理
        return v
    return v


def values_match(a: Any, b: Any, tol: float = 0.01) -> bool:
    """比较两个值是否匹配（支持浮点容差和类型归一化）"""
    a, b = _normalize_cell(a), _normalize_cell(b)
    if a is None and b is None:
        return True
    if a is None or b is None:
        return False
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        if isinstance(a, float) or isinstance(b, float):
            return abs(float(a) - float(b)) < tol
        return a == b
    return a == b


def rows_match(excel_row: list, sqlite_row: tuple, tol: float = 0.01) -> bool:
    """比较 ExcelMCP 行和 SQLite 行"""
    if len(excel_row) != len(sqlite_row):
        return False
    return all(values_match(a, b, tol) for a, b in zip(excel_row, sqlite_row))


def assert_query_match(
    excel_result: dict,
    sqlite_rows: list[tuple],
    sql: str,
    tol: float = 0.01,
):
    """断言 ExcelMCP 查询结果与 SQLite 结果一致"""
    assert excel_result["success"], (
        f"ExcelMCP query failed for SQL: {sql}\n"
        f"Message: {excel_result.get('message', 'N/A')}"
    )

    excel_data = excel_result["data"]
    if len(excel_data) <= 1:
        # ExcelMCP 返回只有表头（或空），检查 SQLite 也为空
        assert len(sqlite_rows) == 0, (
            f"ExcelMCP returned {len(excel_data)} rows (header only), "
            f"but SQLite returned {len(sqlite_rows)} rows\n"
            f"SQL: {sql}\n"
            f"SQLite rows: {sqlite_rows}"
        )
        return

    excel_rows = excel_data[1:]  # 跳过表头
    assert len(excel_rows) == len(sqlite_rows), (
        f"Row count mismatch for SQL: {sql}\n"
        f"ExcelMCP: {len(excel_rows)} rows\n"
        f"SQLite: {len(sqlite_rows)} rows\n"
        f"ExcelMCP data: {excel_rows}\n"
        f"SQLite data: {sqlite_rows}"
    )

    mismatches = []
    for i, (erow, srow) in enumerate(zip(excel_rows, sqlite_rows)):
        if not rows_match(erow, srow, tol):
            mismatches.append((i, erow, srow))

    assert len(mismatches) == 0, (
        f"Data mismatch for SQL: {sql}\n"
        f"Mismatched rows ({len(mismatches)}/{len(excel_rows)}):\n"
        + "\n".join(
            f"  Row {idx}: ExcelMCP={erow} vs SQLite={srow}"
            for idx, erow, srow in mismatches
        )
    )


def assert_affected_rows_match(
    excel_result: dict,
    sqlite_affected: int,
    sql: str,
):
    """断言 ExcelMCP affected_rows 与 SQLite rowcount 一致"""
    assert excel_result["success"], (
        f"ExcelMCP write failed for SQL: {sql}\n"
        f"Message: {excel_result.get('message', 'N/A')}"
    )
    excel_affected = excel_result.get("affected_rows", -1)
    assert excel_affected == sqlite_affected, (
        f"affected_rows mismatch for SQL: {sql}\n"
        f"ExcelMCP: {excel_affected}\n"
        f"SQLite: {sqlite_affected}"
    )


# ============================================================
# Pytest Fixtures
# ============================================================


@pytest.fixture(autouse=True)
def _reset_engine():
    """每个测试前重置 ExcelMCP 引擎缓存"""
    reset_excel_engine()
    yield
    reset_excel_engine()


@pytest.fixture
def adv_tmp_dir(tmp_path) -> Path:
    """对抗测试专用临时目录"""
    return tmp_path


@pytest.fixture
def seed_excel(adv_tmp_dir) -> str:
    """创建种子 Excel 文件（每个测试独立副本）"""
    return _save_wb(_make_seed_wb(), adv_tmp_dir, "seed.xlsx")


@pytest.fixture
def oracle_conn() -> sqlite3.Connection:
    """创建 SQLite oracle 连接（每个测试独立内存 DB）"""
    conn = _create_oracle_db()
    yield conn
    conn.close()


@pytest.fixture
def both(seed_excel: str, oracle_conn: sqlite3.Connection):
    """同时提供 seed_excel 路径和 oracle 连接的便捷 fixture"""
    return seed_excel, oracle_conn


# ============================================================
# 评分收集器（用于 generate_score_report）
# ============================================================


class ScoreCollector:
    """收集对抗测试结果，生成评分报告"""

    def __init__(self):
        self.results: list[dict] = []

    def record(
        self,
        category: str,
        sub_category: str,
        sql: str,
        passed: bool,
        expected: Any = None,
        actual: Any = None,
        error_msg: str = "",
    ):
        self.results.append({
            "category": category,
            "sub_category": sub_category,
            "sql": sql,
            "passed": passed,
            "expected": expected,
            "actual": actual,
            "error_msg": error_msg,
        })

    def summary(self) -> dict:
        total = len(self.results)
        passed = sum(1 for r in self.results if r["passed"])
        failed_list = [r for r in self.results if not r["passed"]]

        # 按操作类型分类
        by_category: dict[str, dict] = {}
        for r in self.results:
            cat = r["category"]
            if cat not in by_category:
                by_category[cat] = {"total": 0, "passed": 0, "failed": []}
            by_category[cat]["total"] += 1
            if r["passed"]:
                by_category[cat]["passed"] += 1
            else:
                by_category[cat]["failed"].append(r)

        return {
            "total": total,
            "passed": passed,
            "accuracy": round(passed / total * 100, 1) if total > 0 else 0,
            "by_category": {
                k: {
                    "total": v["total"],
                    "passed": v["passed"],
                    "accuracy": round(v["passed"] / v["total"] * 100, 1) if v["total"] > 0 else 0,
                    "failed_count": len(v["failed"]),
                }
                for k, v in by_category.items()
            },
            "failed_count": len(failed_list),
            "failed_cases": [
                {
                    "category": r["category"],
                    "sub_category": r["sub_category"],
                    "sql": r["sql"],
                    "expected": r["expected"],
                    "actual": r["actual"],
                    "error_msg": r["error_msg"],
                }
                for r in failed_list
            ],
        }


@pytest.fixture
def score_collector():
    """评分收集器 fixture"""
    return ScoreCollector()


__all__ = [
    "SEED_TABLE",
    "SEED_COLUMNS",
    "SEED_DATA",
    "query_excel",
    "query_sqlite",
    "write_excel",
    "write_sqlite",
    "reset_excel_engine",
    "assert_query_match",
    "assert_affected_rows_match",
    "values_match",
    "rows_match",
    "ScoreCollector",
    "seed_excel",
    "oracle_conn",
    "both",
    "score_collector",
]
